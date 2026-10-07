"""The readers, the filters, the Python API and the command line.

Run with:  uv run python -m unittest discover -s tests

The input files are small synthetic ones written by the tests themselves, in the shapes the
publishers use -- real papers cannot be committed, most of them being licensed for mining, not
redistribution. The model is a fake that answers instantly, so nothing here needs a key or a
network, and the whole pipeline (add, extract, judge, export) still runs end to end.
"""
import contextlib
import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace

os.environ["TOOLKIT_SEED_DEMO"] = "0"

import file2records as fr                      # noqa: E402
from file2records import cli, filters, llm, readers, storage  # noqa: E402

JATS = """<?xml version="1.0" encoding="UTF-8"?>
<article xmlns:xlink="http://www.w3.org/1999/xlink" article-type="research-article">
<front>
  <journal-meta><journal-title-group><journal-title>RSC Advances</journal-title></journal-title-group></journal-meta>
  <article-meta>
    <article-id pub-id-type="doi">10.1039/d0ra00001a</article-id>
    <title-group><article-title>Zinc acetate catalysed <italic>PET</italic> glycolysis</article-title></title-group>
    <pub-date><year>2021</year></pub-date>
    <abstract><p>We depolymerise PET with ethylene glycol.</p></abstract>
  </article-meta>
</front>
<body>
  <sec><title>Results</title>
    <p>Glycolysis of PET over Zn(OAc)<sub>2</sub> gave BHET in high yield <xref ref-type="bibr">[3]</xref>.</p>
    <table-wrap id="t1"><label>Table 1</label><caption><p>Glycolysis results</p></caption>
      <table><thead><tr><th>Catalyst</th><th>T (°C)</th><th>Yield (%)</th></tr></thead>
      <tbody><tr><td>Zn(OAc)<sub>2</sub></td><td>196</td><td>85</td></tr>
             <tr><td>none</td><td>196</td><td>2</td></tr></tbody></table>
      <table-wrap-foot><fn><p>Conditions: PET 2 g, EG 22 g, 3 h.</p></fn></table-wrap-foot>
    </table-wrap>
  </sec>
</body>
<back><ref-list><ref><mixed-citation>Someone et al., 2001.</mixed-citation></ref></ref-list></back>
</article>"""

ELSEVIER = """<?xml version="1.0" encoding="UTF-8"?>
<full-text-retrieval-response xmlns="http://www.elsevier.com/xml/svapi/article/dtd"
  xmlns:prism="http://prismstandard.org/namespaces/basic/2.0/" xmlns:dc="http://purl.org/dc/elements/1.1/"
  xmlns:ce="http://www.elsevier.com/xml/common/dtd" xmlns:xocs="http://www.elsevier.com/xml/xocs/dtd"
  xmlns:ja="http://www.elsevier.com/xml/ja/dtd" xmlns:cals="http://www.elsevier.com/xml/common/cals/dtd">
<coredata>
  <prism:doi>10.1016/j.polymdegradstab.2020.109000</prism:doi>
  <dc:title>Ionic liquid glycolysis of PET</dc:title>
  <prism:publicationName>Polymer Degradation and Stability</prism:publicationName>
  <prism:coverDate>2020-03-01</prism:coverDate>
  <dc:description>Abstract text about [Bmim]Cl.</dc:description>
</coredata>
<originalText><xocs:doc><xocs:serial-item><ja:article><ja:body><ce:sections>
  <ce:section><ce:section-title>Experimental</ce:section-title>
    <ce:para>PET (5 g) was heated with [P<ce:inf>66614</ce:inf>]Cl at 180 °C <ce:cross-ref>[4]</ce:cross-ref>.</ce:para>
    <ce:table><ce:label>Table 2</ce:label><ce:caption><ce:simple-para>Effect of temperature</ce:simple-para></ce:caption>
      <cals:tgroup cols="2"><cals:thead><cals:row><ce:entry>T (°C)</ce:entry><ce:entry>Conversion (%)</ce:entry></cals:row></cals:thead>
      <cals:tbody><cals:row><ce:entry>180</ce:entry><ce:entry>100</ce:entry></cals:row></cals:tbody></cals:tgroup>
      <ce:table-footnote><ce:note-para>Reaction time 4 h.</ce:note-para></ce:table-footnote>
    </ce:table>
  </ce:section>
</ce:sections></ja:body></ja:article></xocs:serial-item></xocs:doc></originalText>
</full-text-retrieval-response>"""

HTML = """<html><head><title>Page</title>
<meta name="citation_title" content="Methanolysis of PET">
<meta name="citation_doi" content="doi:10.1021/acs.iecr.1c00001">
<script>var tracking = 1;</script></head>
<body><nav>Journals | Login</nav>
<h2>Results</h2><p>Methanolysis at 200 °C gave DMT.</p>
<table><caption>Table 1. Runs</caption><tr><th>Run</th><th>DMT (%)</th></tr><tr><td>1</td><td>92</td></tr></table>
</body></html>"""


def fake_completion(params, messages, **kwargs):
    """Answers like a model would, instantly: records for an extraction, verdicts for a judge."""
    system = messages[0]["content"]
    if '"verdicts"' in system:
        reply = {"verdicts": [{"record_index": 0, "critique": "Matches Table 1.", "bad_fields": [],
                               "verdict": "correct", "fixes": [], "drop_record": False}]}
    else:
        reply = {"records": [{"catalyst": "Zn(OAc)2", "temperature_c": 196, "yield_percent": 85,
                              "source_chunk_ids": ["c4"]}]}
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(reply)))],
                           usage=SimpleNamespace(prompt_tokens=100, completion_tokens=20))


class Files:
    """A temporary folder of paper files in every supported shape."""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="f2r-files-"))
        (self.dir / "jats-no-doctype.xml").write_text(JATS, encoding="utf-8")
        (self.dir / "elsevier.xml").write_text(ELSEVIER, encoding="utf-8")
        (self.dir / "page.html").write_text(HTML, encoding="utf-8")
        (self.dir / "notes.md").write_text("# Notes\n\nPET glycolysis at 190 °C.\n\n"
                                           "| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8")
        import docx
        document = docx.Document()
        document.add_heading("Thesis chapter", level=1)
        document.add_paragraph("Hydrolysis of PET gave TPA. doi:10.1002/app.12345")
        table = document.add_table(rows=2, cols=2)
        for (r, c), text in {(0, 0): "Base", (0, 1): "TPA (%)", (1, 0): "NaOH", (1, 1): "97"}.items():
            table.cell(r, c).text = text
        document.save(self.dir / "chapter.docx")

    def __getitem__(self, name):
        return self.dir / name


class ReaderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.files = Files()

    def parse(self, name):
        return readers.parse(self.files[name], "p")

    def test_jats_without_a_doctype_is_still_recognised(self):
        """Europe PMC serves JATS with no DOCTYPE line; docling rejected exactly such a file."""
        chunks, meta, fmt = self.parse("jats-no-doctype.xml")
        self.assertEqual(fmt, "jats")
        self.assertEqual(meta["doi"], "10.1039/d0ra00001a")
        self.assertEqual(meta["title"], "Zinc acetate catalysed PET glycolysis")
        self.assertEqual(meta["year"], "2021")
        text = "\n\n".join(c["text"] for c in chunks)
        self.assertIn("Zn(OAc)2", text, "inline markup must dissolve into the word")
        self.assertNotIn("Someone et al.", text, "the reference list is not paper content")

    def test_a_table_keeps_its_rows_and_its_footnote(self):
        for name, row, footnote in (("jats-no-doctype.xml", "| Zn(OAc)2 | 196 | 85 |", "PET 2 g, EG 22 g"),
                                    ("elsevier.xml", "| 180 | 100 |", "Reaction time 4 h")):
            table = next(c["text"] for c in self.parse(name)[0] if "| --- |" in c["text"])
            self.assertIn(row, table, name)
            self.assertIn(footnote, table, f"{name}: the conditions every row shares were lost")

    def test_elsevier_xml(self):
        chunks, meta, fmt = self.parse("elsevier.xml")
        self.assertEqual(fmt, "elsevier")
        self.assertEqual(meta, {"doi": "10.1016/j.polymdegradstab.2020.109000",
                                "title": "Ionic liquid glycolysis of PET",
                                "journal": "Polymer Degradation and Stability", "year": "2020"})
        self.assertTrue(any("[P66614]Cl" in c["text"] for c in chunks))
        self.assertTrue(any(c["text"] == "## Experimental" for c in chunks))

    def test_html_reads_the_article_not_the_page_furniture(self):
        chunks, meta, fmt = self.parse("page.html")
        self.assertEqual((fmt, meta["doi"], meta["title"]),
                         ("html", "10.1021/acs.iecr.1c00001", "Methanolysis of PET"))
        text = " ".join(c["text"] for c in chunks)
        self.assertIn("| 1 | 92 |", text)
        self.assertNotIn("tracking", text)
        self.assertNotIn("Login", text)

    def test_word_and_markdown(self):
        chunks, meta, fmt = self.parse("chapter.docx")
        self.assertEqual((fmt, meta["doi"]), ("docx", "10.1002/app.12345"))
        self.assertIn("| NaOH | 97 |", " ".join(c["text"] for c in chunks))
        chunks, _, fmt = self.parse("notes.md")
        self.assertEqual(fmt, "markdown")
        self.assertTrue(any(c["text"].startswith("| a | b |") for c in chunks))

    def test_chunk_ids_are_stable(self):
        self.assertEqual(self.parse("elsevier.xml")[0], self.parse("elsevier.xml")[0])

    def test_wrong_or_unsupported_files_say_what_is_wrong(self):
        bad = self.files.dir / "fake.pdf"
        bad.write_text("not a pdf")
        with self.assertRaisesRegex(readers.UnsupportedFile, "not a PDF"):
            readers.detect(bad)
        odd = self.files.dir / "data.xlsx"
        odd.write_bytes(b"PK")
        with self.assertRaisesRegex(readers.UnsupportedFile, "unsupported file type"):
            readers.detect(odd)


class ProjectTests(unittest.TestCase):
    """The Python API end to end, with a fake model."""

    def setUp(self):
        self.files = Files()
        self.dir = Path(tempfile.mkdtemp(prefix="f2r-project-"))
        self.previous = storage.WORKSPACE
        self.project = fr.Project(self.dir)
        self.real_complete, llm.complete = llm.complete, fake_completion

    def tearDown(self):
        llm.complete = self.real_complete
        if self.previous != Path():
            storage.use(self.previous)
        shutil.rmtree(self.dir, ignore_errors=True)

    def configure(self):
        self.project.schema = [{"name": "catalyst", "type": "string", "description": "as written"},
                               {"name": "temperature_c", "type": "number"},
                               {"name": "yield_percent", "type": "number"}]
        self.project.prompt = "Extract every glycolysis experiment."
        self.project.rubric = "Check each record against the paper."

    def test_add_extract_judge_export(self):
        self.configure()
        added = self.project.add(self.files.dir)
        self.assertEqual(sum(1 for r in added if r.get("error")), 0, added)
        self.assertEqual(len(self.project.papers()), 5)

        done = self.project.extract("gpt-4o-mini", only="glycoly[sz]is")
        self.assertEqual(len(done), 3, "jats, elsevier and the markdown note mention glycolysis")
        self.assertEqual(self.project.extract("gpt-4o-mini", only="glycoly[sz]is"), [],
                         "a second run must skip what is done")
        self.assertEqual(len(self.project.extract("gpt-4o-mini", only="glycoly[sz]is", redo=True)), 3)

        judged = self.project.judge("gpt-4o-mini")
        self.assertEqual(len(judged), 3, "only extracted papers are judged")

        rows = self.project.records(only=r"Zinc acetate")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["doi"], "10.1039/d0ra00001a")
        self.assertEqual(rows[0]["judge_verdict"], "correct")
        self.assertTrue(rows[0]["source_chunk_ids"], "the c4 citation must resolve to a chunk id")

        csv_path = self.project.export(self.dir / "out.csv", exclude="Elsevier|elsevier|Ionic")
        self.assertNotIn("polymdegradstab", csv_path.read_text())

    def test_a_bundle_carries_no_paper_text_unless_asked(self):
        self.configure()
        self.project.add(self.files["jats-no-doctype.xml"])
        self.project.extract("gpt-4o-mini")
        for include_text in (False, True):
            path = self.project.export(self.dir / "b.zip", include_text=include_text)
            with zipfile.ZipFile(path) as z:
                parsed = json.loads(z.read(next(n for n in z.namelist() if n.startswith("parsed/"))))
                manifest = json.loads(z.read("manifest.json"))
            has_text = any("text" in c for c in parsed["chunks"])
            self.assertEqual(has_text, include_text)
            self.assertEqual(manifest["includes_paper_text"], include_text)
            self.assertTrue(all("id" in c for c in parsed["chunks"]), "chunk ids always travel")

    def test_stages_refuse_with_reasons_before_calling_anything(self):
        with self.assertRaisesRegex(RuntimeError, "Define the fields"):
            self.project.extract("gpt-4o-mini")
        self.configure()
        self.assertTrue(any("No model yet" in m for m in self.project.check("extract")))

    def test_a_model_is_found_however_the_service_spells_it(self):
        """KI:connect calls one model "Mistral Small 4 119b" in its chat menu,
        "mistral-small-4-119b-2603" on its overview page and "mistralai-mistral-small-4-119b"
        in its API. Any of them works, the "openai/" prefix is added, embeddings are dropped."""
        served = ["mistralai-mistral-small-4-119b", "Qwen 3.8 27B", "gpt-oss-120b",
                  "gpt-6-luna", "gpt-6-sol", "e5-mistral-7b-instruct", "qwen3-embedding-8b"]
        real = llm.list_models
        llm.list_models = lambda endpoint, key: [i for i in served if not llm.NOT_CHAT.search(i)]
        endpoint = "https://chat.kiconnect.nrw/api/v1"
        try:
            for spelling, api_name in [("Mistral Small 4 119b", "mistralai-mistral-small-4-119b"),
                                       ("mistral-small-4-119b-2603", "mistralai-mistral-small-4-119b"),
                                       ("mistral", "mistralai-mistral-small-4-119b"),
                                       ("OpenAI GPT OSS 120b", "gpt-oss-120b"),
                                       ("qwen3.8-27b", "Qwen 3.8 27B")]:
                self.assertEqual(fr.connect("id:secret", endpoint, spelling),
                                 {"model": "openai/" + api_name, "api_base": endpoint + "/",
                                  "api_key": "id:secret"}, spelling)
            with self.assertRaisesRegex(RuntimeError, "Choose a model. This service offers: "
                                                      "mistralai-mistral-small-4-119b, Qwen"):
                fr.connect("k", endpoint)
            with self.assertRaisesRegex(RuntimeError, "matches several"):
                fr.connect("k", endpoint, "gpt-6")
            with self.assertRaisesRegex(RuntimeError, "is not one of"):
                fr.connect("k", endpoint, "e5-mistral")     # an embedding model, not offered
            self.assertEqual(fr.connect("sk-x", model="gpt-4o-mini"),
                             {"model": "gpt-4o-mini", "api_key": "sk-x"})
        finally:
            llm.list_models = real

    def test_schema_from_a_pydantic_model(self):
        from pydantic import BaseModel, Field

        class Run(BaseModel):
            catalyst: str | None = Field(None, description="as written")
            yield_percent: float | None = None
            repeats: int = 1

        self.project.schema = Run
        self.assertEqual(self.project.schema, [
            {"name": "catalyst", "type": "string", "description": "as written"},
            {"name": "yield_percent", "type": "number", "description": ""},
            {"name": "repeats", "type": "integer", "description": ""}])

    def test_search_and_select(self):
        self.project.add(self.files.dir)
        hits = self.project.search(r"glycoly[sz]is")
        self.assertEqual({h["filename"] for h in hits},
                         {"jats-no-doctype.xml", "elsevier.xml", "notes.md"})
        self.assertTrue(all(s["match"].lower().startswith("glycoly") for h in hits for s in h["snippets"]))
        both = self.project.select(only="PET", exclude="glycoly")
        self.assertEqual(len(both), 2, "the HTML and the Word file mention PET but not glycolysis")
        with self.assertRaisesRegex(ValueError, "Not a valid regular expression"):
            filters.compile_pattern("(unclosed")


class CliTests(unittest.TestCase):
    def setUp(self):
        self.files = Files()
        self.dir = Path(tempfile.mkdtemp(prefix="f2r-cli-"))
        self.previous = storage.WORKSPACE

    def tearDown(self):
        if self.previous != Path():
            storage.use(self.previous)
        shutil.rmtree(self.dir, ignore_errors=True)

    def run_cli(self, *argv):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = cli.main([str(a) for a in argv])
        return code, out.getvalue() + err.getvalue()

    def test_add_search_check_export(self):
        code, out = self.run_cli("add", self.dir, self.files.dir)
        self.assertEqual(code, 0, out)
        self.assertIn("5 added, 0 failed", out)
        self.assertIn("[jats]", out)

        code, out = self.run_cli("search", self.dir, "glycoly[sz]is")
        self.assertIn("3 papers", out)

        code, out = self.run_cli("check", self.dir)
        self.assertEqual(code, 1, "nothing is configured, so extraction is not ready")
        self.assertIn("No model yet", out)

        code, out = self.run_cli("export", self.dir, self.dir / "out.json")
        self.assertEqual((code, json.loads((self.dir / "out.json").read_text())), (0, []))

    def test_a_bad_regex_is_an_error_message_not_a_traceback(self):
        self.run_cli("add", self.dir, self.files["notes.md"])
        code, out = self.run_cli("search", self.dir, "(unclosed")
        self.assertEqual(code, 2)
        self.assertIn("error: Not a valid regular expression", out)


if __name__ == "__main__":
    unittest.main()
